from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Guia de juego Rural Murdoku.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, size=None, bold=None, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    style_run(run, 10.5)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(17)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text)
    style_run(run, 15, True, (0, 0, 0))
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    style_run(run, 10.5)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)

for style_name in ("Title", "Heading 1", "Heading 2"):
    style = doc.styles[style_name]
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
run = title.add_run("Guía de juego Rural Murdoku")
style_run(run, 27, True, (0, 0, 0))

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(24)
run = subtitle.add_run("Cómo participar en la investigación del Compás Dorado")
style_run(run, 13, False, (70, 70, 70))

intro_table = doc.add_table(rows=1, cols=1)
intro_table.autofit = False
cell = intro_table.cell(0, 0)
set_cell_shading(cell, "EAF0F8")
set_cell_margins(cell, top=170, start=210, bottom=170, end=210)
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
run = p.add_run(
    "Rural Murdoku es un juego de investigación por equipos. La aplicación reúne las pistas, "
    "las misiones y los recursos de cada grupo para que la partida avance de forma ordenada."
)
style_run(run, 11)

add_heading(doc, "La partida")
add_body(
    doc,
    "Al comenzar, cada participante recibe un personaje, un equipo y una posición secreta en el tablero. "
    "Nadie ve toda la información: cada jugador cuenta únicamente con lo que su equipo ha descubierto."
)
add_body(
    doc,
    "El propósito general es reunir información útil sobre el robo del Compás Dorado. Hablar, contrastar "
    "pistas y decidir en qué invertir los recursos del equipo forma parte esencial de la partida."
)

add_heading(doc, "Pistas y misiones")
add_body(
    doc,
    "Cada equipo empieza con algunas pistas básicas sobre sus propios integrantes. Después, las misiones "
    "secundarias permiten conseguir nuevas pistas relacionadas con personajes de otros equipos."
)
add_bullet(doc, "Cuando una misión secundaria se completa, la aplicación registra el resultado y añade la pista correspondiente al listado del equipo.")
add_bullet(doc, "En ocasiones una pista puede perderse. No desaparece para siempre: queda disponible para recuperarla mediante los recursos del equipo.")
add_bullet(doc, "El administrador decide cuándo se habilitan las compras y cuándo se liberan las pistas comunes relacionadas con el Compás.")

add_heading(doc, "Cuadrantes y ubicaciones")
add_body(
    doc,
    "Los cuadrantes ayudan a situar la investigación en el tablero. Un equipo puede obtener nuevos cuadrantes "
    "y, cuando dispone de uno, decidir si compra también las ubicaciones que contiene."
)
add_body(
    doc,
    "Algunos cuadrantes se consiguen mediante pruebas o juegos fuera de la aplicación. Cuando un equipo supera "
    "esa prueba, el administrador habilita el cuadrante para que el equipo pueda utilizarlo en la app."
)

add_heading(doc, "Roles secretos")
add_body(
    doc,
    "La mayoría de participantes juega como investigador. Sin embargo, existen un ladrón y dos cómplices. "
    "Sus roles solo aparecen en su información secreta y comparten una palabra para poder reconocerse sin llamar la atención."
)
add_body(
    doc,
    "Durante ciertas rondas, estos roles pueden alterar la investigación. Sus acciones pueden impedir que un "
    "jugador reciba una pista o afectar a los recursos generales. El resto de jugadores deberá estar atento a "
    "lo que ocurre y seguir construyendo sus deducciones."
)

add_heading(doc, "Monedas y decisiones de equipo")
add_body(
    doc,
    "Cada equipo dispone de monedas. Se utilizan para recuperar pistas perdidas, obtener cuadrantes o descubrir "
    "ubicaciones. Antes de gastar, conviene hablarlo con el equipo: no todas las compras aportan la misma información "
    "en el mismo momento."
)

add_heading(doc, "Eventos de grupo")
add_body(
    doc,
    "A lo largo de la partida habrá dos eventos grupales programados a horas determinadas de forma aleatoria. "
    "La aplicación avisará a los jugadores mediante una notificación."
)
add_bullet(doc, "Reunión en una estancia: todos deberán acudir a la sala indicada dentro de la casa.")
add_bullet(doc, "Achupé: todos deberán sentarse; el último en hacerlo asumirá las consecuencias que se indiquen durante el juego.")

add_heading(doc, "Uso de la aplicación")
add_bullet(doc, "Consulta con frecuencia el listado de pistas de tu equipo y tu misión secundaria actual.")
add_bullet(doc, "Mantén instaladas y activadas las notificaciones para no perder los eventos grupales.")
add_bullet(doc, "No enseñes a otros jugadores la información secreta de tu personaje.")
add_bullet(doc, "Si una acción no está disponible, espera a que el administrador habilite la siguiente ronda de compras.")

add_body(
    doc,
    "La aplicación es una herramienta para organizar el juego, pero la investigación ocurre entre las personas. "
    "Observad, hablad, dudid y disfrutad de la partida."
)

doc.core_properties.title = "Guía de juego Rural Murdoku"
doc.core_properties.subject = "Mecánica general para participantes"
doc.core_properties.author = "Rural Murdoku"
doc.save(OUTPUT)
print(OUTPUT)
